#!/usr/bin/env python3
"""Replace CDS Studies input SQL with Memgraph Cypher (TC02 / phs000720 patterns)."""

from __future__ import annotations

import io
import re
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

REPO = Path("/Users/sohilz2/Automation/Commons_Automation")
STUDIES = REPO / "Test Cases/CDS_TestCases/Studies"
INPUTS = REPO / "InputFiles/CDS"
SKIP_STUDIES = {"phs000720"}
XLSX_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

# SQL alias.column -> (cypher_alias, cypher_prop, mode)
COL_MAP = {
    ("s", "phs_accession"): ("s", "phs_accession", "eq"),
    ("s", "study_name"): ("s", "study_name", "eq"),
    ("s", "study_data_types"): ("s", "study_data_types", "contains_pipe"),
    ("sp", "gender"): ("p", "sex", "eq"),
    ("sp", "sex"): ("p", "sex", "eq"),
    ("smp", "sample_tumor_status"): ("samp", "sample_tumor_status", "eq"),
    ("smp", "sample_type"): ("samp", "sample_type", "eq"),
    ("f", "file_type"): ("f", "file_type", "eq_ci"),
    ("f1", "file_type"): ("f", "file_type", "eq_ci"),
    ("f", "experimental_strategy_and_data_subtypes"): (
        "f",
        "experimental_strategy_and_data_subtypes",
        "contains_pipe",
    ),
    ("f1", "experimental_strategy_and_data_subtypes"): (
        "f",
        "experimental_strategy_and_data_subtypes",
        "contains_pipe",
    ),
    ("d", "primary_diagnosis"): ("d", "primary_diagnosis", "eq"),
    ("gi", "library_strategy"): ("gi", "library_strategy", "eq"),
    ("gi", "library_layout"): ("gi", "library_layout", "eq"),
    ("gi", "library_selection"): ("gi", "library_selection", "eq"),
    ("gi", "library_source"): ("gi", "library_source", "eq"),
    ("gi", "platform"): ("gi", "platform", "eq"),
    ("gi", "instrument_model"): ("gi", "instrument_model", "eq"),
    ("gi", "reference_genome_assembly"): ("gi", "reference_genome_assembly", "eq"),
}

VALUE_REMAP = {
    ("gi", "library_source", "Transcriptomic"): "Transcriptome",
    ("gi", "library_strategy", "ATAC-seq"): "ATAC-Seq",
    ("gi", "library_strategy", "OTHER"): "Other",
    ("gi", "library_selection", "other"): "Other",
    ("gi", "library_selection", "PolyA"): "Poly-A Enriched Genomic Library",
    ("smp", "sample_type", "Skin"): "Skin, NOS",
    ("d", "primary_diagnosis", "Yolk sac tumor"): "Yolk Sac Tumor",
}

# Equality values that should become CONTAINS (DB labels are more specific)
CONTAINS_VALUES = {
    ("gi", "library_source", "Genomic"),
}


def cq(val: str) -> str:
    return '"' + val.replace("\\", "\\\\").replace('"', '\\"') + '"'


def extract_sql_string(sql: str, open_quote_idx: int) -> tuple[str, int]:
    """Parse SQL string allowing unescaped apostrophes inside values."""
    i = open_quote_idx + 1
    while i < len(sql):
        if sql[i] == "'":
            if i + 1 < len(sql) and sql[i + 1] == "'":
                i += 2
                continue
            rest = sql[i + 1 :].lstrip()
            upper = rest.upper()
            if rest == "" or upper.startswith(
                ("AND", "GROUP", "ORDER", "LIMIT", ")", ";", ",")
            ):
                return sql[open_quote_idx + 1 : i].replace("''", "'"), i + 1
        i += 1
    raise ValueError("unclosed SQL string")


def parse_sql_filters(sql: str) -> list[tuple[str, str, str]]:
    out: list[tuple[str, str, str]] = []
    for m in re.finditer(
        r"(?:^|[\s(])([a-zA-Z0-9_]+)\.([a-zA-Z0-9_]+)\s*=\s*'", sql
    ):
        open_q = m.end() - 1
        try:
            val, _ = extract_sql_string(sql, open_q)
        except ValueError:
            continue
        out.append((m.group(1), m.group(2), val))
    return out


def normalize(raw: list[tuple[str, str, str]]):
    eqs: list[tuple[str, str, str, str]] = []  # alias, prop, val, mode
    contains: list[tuple[str, str, str]] = []
    seen = set()
    for alias, col, val in raw:
        if (alias, col) not in COL_MAP:
            continue
        n_alias, n_prop, mode = COL_MAP[(alias, col)]
        if (alias, col, val) in VALUE_REMAP:
            val = VALUE_REMAP[(alias, col, val)]
        if (alias, col, val) in CONTAINS_VALUES or (
            n_alias,
            n_prop,
            val,
        ) in {
            (COL_MAP[k][0], COL_MAP[k][1], v)
            for (a, c, v) in CONTAINS_VALUES
            for k in [ (a, c) ]
            if k in COL_MAP
        }:
            mode = "contains_pipe"
            # single token contains
        if mode == "contains_pipe":
            tokens = [t.strip() for t in val.split("|") if t.strip()]
            if (alias, col, val) in CONTAINS_VALUES or (n_alias, n_prop, val) in {
                (COL_MAP[(a, c)][0], COL_MAP[(a, c)][1], v) for a, c, v in CONTAINS_VALUES if (a, c) in COL_MAP
            }:
                tokens = [val]
            for tok in tokens:
                if n_prop == "study_data_types" and tok == "Genomic":
                    tok = "Genomics"
                item = ("c", n_alias, n_prop, tok)
                if item in seen:
                    continue
                seen.add(item)
                contains.append((n_alias, n_prop, tok))
        else:
            item = ("e", n_alias, n_prop, val, mode)
            if item in seen:
                continue
            seen.add(item)
            eqs.append((n_alias, n_prop, val, mode))
    return eqs, contains


def where_clause(eqs, contains) -> str:
    parts = []
    for a, p, v, mode in eqs:
        if mode == "eq_ci":
            parts.append(f"toLower({a}.{p}) = toLower({cq(v)})")
        else:
            parts.append(f"{a}.{p} = {cq(v)}")
    parts += [f"{a}.{p} CONTAINS {cq(t)}" for a, p, t in contains]
    return "\n  AND ".join(parts)


def flags_from(eqs, contains):
    aliases = {a for a, _, _, _ in eqs} | {a for a, _, _ in contains}
    return {
        "gi": "gi" in aliases,
        "file": "f" in aliases or "gi" in aliases,
        "samp": "samp" in aliases,
        "diag": "d" in aliases,
    }


def build_base_matches(flags) -> list[str]:
    gi, file_, samp, diag = flags["gi"], flags["file"], flags["samp"], flags["diag"]

    # Sample + genomic/file filters intersect at the participant (720 pattern),
    # not via from_sample — graph often lacks file→sample edges for those combos.
    if gi and samp and diag:
        return [
            "MATCH (samp:sample)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
            "MATCH (gi:genomic_info)-[:of_file]->(f:file)-[:of_participant]->(p)",
            "MATCH (d:diagnosis)-[:of_participant]->(p)",
        ]
    if gi and samp:
        return [
            "MATCH (samp:sample)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
            "MATCH (gi:genomic_info)-[:of_file]->(f:file)-[:of_participant]->(p)",
        ]
    if gi and diag:
        return [
            "MATCH (gi:genomic_info)-[:of_file]->(f:file)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
            "MATCH (d:diagnosis)-[:of_participant]->(p)",
        ]
    if gi:
        return [
            "MATCH (gi:genomic_info)-[:of_file]->(f:file)-[:of_participant]->(p:participant)-[:of_study]->(s:study)"
        ]
    if diag and file_ and samp:
        return [
            "MATCH (d:diagnosis)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
            "MATCH (samp:sample)-[:of_participant]->(p)",
            "MATCH (f:file)-[:of_participant]->(p)",
        ]
    if diag and file_:
        return [
            "MATCH (d:diagnosis)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
            "MATCH (f:file)-[:of_participant]->(p)",
        ]
    if diag and samp:
        return [
            "MATCH (d:diagnosis)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
            "MATCH (samp:sample)-[:of_participant]->(p)",
        ]
    if diag:
        return [
            "MATCH (d:diagnosis)-[:of_participant]->(p:participant)-[:of_study]->(s:study)"
        ]
    if file_ and samp:
        return [
            "MATCH (samp:sample)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
            "MATCH (f:file)-[:of_participant]->(p)",
        ]
    if file_:
        return [
            "MATCH (f:file)-[:of_participant]->(p:participant)-[:of_study]->(s:study)"
        ]
    if samp:
        return [
            "MATCH (samp:sample)-[:of_participant]->(p:participant)-[:of_study]->(s:study)"
        ]
    return ["MATCH (p:participant)-[:of_study]->(s:study)"]


PARTICIPANT_TAIL = """WITH DISTINCT p, s
OPTIONAL MATCH (samp_all:sample)-[:of_participant]->(p)
WITH p, s, samp_all
ORDER BY samp_all.sample_id ASC
WITH p, s, [x IN collect(samp_all.sample_id) WHERE x IS NOT NULL] AS sample_ids
WITH p, s, sample_ids,
  CASE
    WHEN size(sample_ids) > 5
    THEN reduce(acc = "", id IN sample_ids[0..5] | CASE WHEN acc = "" THEN id ELSE acc + ", " + id END) + ", ..."
    ELSE reduce(acc = "", id IN sample_ids | CASE WHEN acc = "" THEN id ELSE acc + ", " + id END)
  END AS samples
RETURN
  p.participant_id AS participant_id,
  s.study_name AS study_name,
  s.phs_accession AS accession,
  p.sex AS sex,
  samples
ORDER BY p.participant_id ASC
LIMIT 100
"""

SAMPLE_TAIL = """RETURN DISTINCT
  s.study_acronym AS study_name,
  s.phs_accession AS accession,
  samp.sample_id AS sample_id,
  COALESCE(samp.sample_name, samp.sample_id) AS sample_name,
  COALESCE(samp.Organization_Name, "Not specified in data") AS organization_name
ORDER BY samp.sample_id ASC
LIMIT 100
"""

FILE_TAIL = """RETURN DISTINCT
  s.study_name AS study_name,
  s.phs_accession AS accession,
  f.file_name AS file_name,
  f.file_id AS file_id,
  f.file_type AS file_type,
  COALESCE(samp.sample_id, "Not Applicable") AS sample_id
ORDER BY f.file_name ASC
LIMIT 100
"""

STAT_TAIL = """RETURN
  count(DISTINCT s) AS Studies,
  count(DISTINCT p) AS Participants,
  count(DISTINCT samp) AS Samples,
  count(DISTINCT f) AS Files
"""


def participants_query(matches, where, flags) -> str:
    # Filter path selects participants; PARTICIPANT_TAIL always reloads ALL samples
    # for those participants (UI Participants-tab behavior).
    lines = list(matches) + [f"WHERE {where}", PARTICIPANT_TAIL]
    return "\n".join(lines)


def samples_query(matches, where, flags) -> str:
    if any("samp:sample" in m for m in matches):
        return "\n".join(list(matches) + [f"WHERE {where}", SAMPLE_TAIL])

    if flags["gi"]:
        ms = [
            "MATCH (gi:genomic_info)-[:of_file]->(f:file)-[:from_sample]->(samp:sample)-[:of_participant]->(p:participant)-[:of_study]->(s:study)"
        ]
        if flags["diag"]:
            ms.append("MATCH (d:diagnosis)-[:of_participant]->(p)")
        return "\n".join(ms + [f"WHERE {where}", SAMPLE_TAIL])

    if flags["file"]:
        if flags["diag"]:
            ms = [
                "MATCH (d:diagnosis)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
                "MATCH (f:file)-[:from_sample]->(samp:sample)-[:of_participant]->(p)",
            ]
        else:
            ms = [
                "MATCH (f:file)-[:from_sample]->(samp:sample)-[:of_participant]->(p:participant)-[:of_study]->(s:study)"
            ]
        return "\n".join(ms + [f"WHERE {where}", SAMPLE_TAIL])

    if flags["diag"]:
        return "\n".join(
            list(matches)
            + [f"WHERE {where}", "MATCH (samp:sample)-[:of_participant]->(p)", SAMPLE_TAIL]
        )

    return "\n".join(
        [
            "MATCH (samp:sample)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
            f"WHERE {where}",
            SAMPLE_TAIL.replace("RETURN DISTINCT", "RETURN"),
        ]
    )


def files_query(matches, where, flags) -> str:
    # Sample filters: files must hang off the *filtered* samples (UI facet scope),
    # not every file on the participant.
    if flags["samp"] and not (flags["file"] or flags["gi"]):
        return "\n".join(
            list(matches)
            + [
                "MATCH (f:file)-[:from_sample]->(samp)",
                f"WHERE {where}",
                FILE_TAIL,
            ]
        )

    if flags["file"] or flags["gi"]:
        lines = list(matches) + [f"WHERE {where}"]
        has_from_sample = any("from_sample" in m for m in matches)
        has_samp = any("samp:sample" in m for m in matches)
        if flags["samp"] and not has_from_sample:
            # Sample + file/genomic filters: keep file predicates, require from_sample link
            lines = list(matches) + [
                "MATCH (f)-[:from_sample]->(samp)",
                f"WHERE {where}",
                FILE_TAIL,
            ]
            return "\n".join(lines)
        if has_samp and not has_from_sample:
            lines.append("OPTIONAL MATCH (f)-[:from_sample]->(samp2:sample)")
            lines.append(
                """RETURN DISTINCT
  s.study_name AS study_name,
  s.phs_accession AS accession,
  f.file_name AS file_name,
  f.file_id AS file_id,
  f.file_type AS file_type,
  COALESCE(samp2.sample_id, samp.sample_id, "Not Applicable") AS sample_id
ORDER BY f.file_name ASC
LIMIT 100
"""
            )
            return "\n".join(lines)
        if not has_samp:
            lines.append("OPTIONAL MATCH (f)-[:from_sample]->(samp:sample)")
        lines.append(FILE_TAIL)
        return "\n".join(lines)

    if flags["diag"]:
        return "\n".join(
            list(matches)
            + [
                "MATCH (f:file)-[:of_participant]->(p)",
                f"WHERE {where}",
                "OPTIONAL MATCH (f)-[:from_sample]->(samp:sample)",
                FILE_TAIL,
            ]
        )

    return "\n".join(
        [
            "MATCH (f:file)-[:of_participant]->(p:participant)-[:of_study]->(s:study)",
            f"WHERE {where}",
            "OPTIONAL MATCH (f)-[:from_sample]->(samp:sample)",
            FILE_TAIL.replace("RETURN DISTINCT", "RETURN"),
        ]
    )


def stat_query(matches, where, flags) -> str:
    lines = list(matches) + [f"WHERE {where}"]
    has_samp = any("samp:sample" in m for m in matches)
    has_file = flags["file"] or flags["gi"] or any(
        "(f:file)" in m or "(f)" in m for m in matches
    )
    if not has_samp:
        # File/GI facet scope: UI Samples = samples linked to the *filtered* files.
        # Participant/study/diagnosis-only: all samples on matched participants.
        if has_file:
            lines.append("OPTIONAL MATCH (f)-[:from_sample]->(samp:sample)")
        else:
            lines.append("OPTIONAL MATCH (samp:sample)-[:of_participant]->(p)")
    if not has_file:
        # Sample filters: count files linked to the *filtered* samples (UI facet scope).
        # Otherwise: all files for the matched participants.
        if has_samp or flags["samp"]:
            lines.append("OPTIONAL MATCH (f:file)-[:from_sample]->(samp)")
        else:
            lines.append("OPTIONAL MATCH (f:file)-[:of_participant]->(p)")
    lines.append(STAT_TAIL)
    return "\n".join(lines)


def generate_from_sql(sql: str, phs_override: str | None = None) -> dict[str, str]:
    eqs, contains = normalize(parse_sql_filters(sql))
    if phs_override:
        eqs = [
            (a, p, phs_override if (a == "s" and p == "phs_accession") else v, m)
            for a, p, v, m in eqs
        ]
        if not any(a == "s" and p == "phs_accession" for a, p, _, _ in eqs):
            eqs.insert(0, ("s", "phs_accession", phs_override, "eq"))
    if not any(a == "s" and p == "phs_accession" for a, p, _, _ in eqs):
        raise ValueError("missing phs_accession")
    where = where_clause(eqs, contains)
    flags = flags_from(eqs, contains)
    matches = build_base_matches(flags)
    return {
        "ParticipantsTab": participants_query(matches, where, flags),
        "StatQuery": stat_query(matches, where, flags),
        "SamplesTab": samples_query(matches, where, flags),
        "FilesTab": files_query(matches, where, flags),
    }


def phs_from_filename(path: Path) -> str | None:
    # Study accessions are always phs + 6 digits (e.g. phs004225).
    # Filenames like TC01_CDS_phs00422510_... must not become phs00422510.
    m = re.search(r"(phs\d{6})", path.stem, re.I)
    return m.group(1).lower() if m else None


def index_inputs() -> dict[str, Path]:
    by_stem: dict[str, list[Path]] = {}
    for p in INPUTS.rglob("*.xlsx"):
        if p.name.startswith("~$") or "Copy" in p.name:
            continue
        if "old set" in str(p).lower():
            continue
        by_stem.setdefault(p.stem, []).append(p)
    chosen = {}
    for stem, cands in by_stem.items():
        preferred = sorted(
            cands,
            key=lambda x: (
                "Single Filter" in str(x),
                x.parent.name.startswith("phs"),
                len(str(x)),
            ),
        )
        chosen[stem] = preferred[0]
    return chosen


def read_participants_sql_from_xlsx_bytes(data: bytes) -> str | None:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        ss = []
        if "xl/sharedStrings.xml" in names:
            ss = [
                "".join(si.itertext())
                for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall(
                    f"{XLSX_NS}si"
                )
            ]
        sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
        header = None
        for row in sheet.findall(f"{XLSX_NS}sheetData/{XLSX_NS}row"):
            vals = []
            for c in row.findall(f"{XLSX_NS}c"):
                t = c.get("t")
                v = c.find(f"{XLSX_NS}v")
                is_el = c.find(f"{XLSX_NS}is")
                if t == "inlineStr" and is_el is not None:
                    vals.append("".join(is_el.itertext()))
                elif v is not None:
                    vals.append(ss[int(v.text)] if t == "s" else (v.text or ""))
                else:
                    vals.append("")
            if not vals:
                continue
            if header is None:
                header = vals
                continue
            if vals[0] == "ParticipantsTab":
                return vals[1] if len(vals) > 1 else None
    return None


def git_show_sql(path: Path) -> str | None:
    rel = path.resolve().relative_to(REPO.resolve())
    for rev in ("HEAD", "HEAD~1"):
        try:
            data = subprocess.check_output(
                ["git", "show", f"{rev}:{rel.as_posix()}"],
                cwd=str(REPO),
                stderr=subprocess.DEVNULL,
            )
        except subprocess.CalledProcessError:
            continue
        try:
            sql = read_participants_sql_from_xlsx_bytes(data)
        except Exception:
            continue
        if sql and not str(sql).lstrip().upper().startswith("MATCH"):
            return sql
    return None


def update_file(path: Path, dry_run: bool = False, force_from_git: bool = False) -> str:
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        tab_i = headers.index("TabName")
        query_i = headers.index("TabQuery")
        stat_i = headers.index("StatQuery")
    except ValueError:
        return f"SKIP bad headers {headers}"

    part_sql = None
    for row in ws.iter_rows(min_row=2):
        if row[tab_i].value == "ParticipantsTab":
            part_sql = row[query_i].value or ""
            break
    if not part_sql:
        return "SKIP no ParticipantsTab"

    if str(part_sql).lstrip().upper().startswith("MATCH"):
        if not force_from_git:
            return "SKIP already cypher"
        git_sql = git_show_sql(path)
        if not git_sql:
            return "SKIP already cypher (no git SQL)"
        part_sql = git_sql

    try:
        gen = generate_from_sql(str(part_sql), phs_override=phs_from_filename(path))
    except Exception as e:
        return f"FAIL {e}"

    for row in ws.iter_rows(min_row=2):
        tab = row[tab_i].value
        if tab == "ParticipantsTab":
            row[query_i].value = gen["ParticipantsTab"]
            row[stat_i].value = gen["StatQuery"]
        elif tab == "SamplesTab":
            row[query_i].value = gen["SamplesTab"]
            if row[stat_i].value not in (None, ""):
                row[stat_i].value = ""
        elif tab == "FilesTab":
            row[query_i].value = gen["FilesTab"]
            if row[stat_i].value not in (None, ""):
                row[stat_i].value = ""

    if not dry_run:
        wb.save(path)
    return "UPDATED"


def main():
    dry = "--dry-run" in sys.argv
    force = "--from-git" in sys.argv
    inputs = index_inputs()
    targets: list[Path] = []
    missing: list[str] = []

    for tc in sorted(STUDIES.rglob("*.tc")):
        study = tc.parent.name
        if study in SKIP_STUDIES or study == "StaticData":
            continue
        if tc.stem.startswith("GC_") or tc.stem.endswith(".xlsx"):
            continue
        if tc.stem not in inputs:
            missing.append(tc.stem)
            continue
        targets.append(inputs[tc.stem])

    seen = set()
    uniq = []
    for p in targets:
        rp = p.resolve()
        if rp in seen:
            continue
        seen.add(rp)
        uniq.append(p)

    counts = {"UPDATED": 0, "SKIP already cypher": 0, "other": 0}
    failures = []
    for p in uniq:
        status = update_file(p, dry_run=dry, force_from_git=force)
        if status == "UPDATED":
            counts["UPDATED"] += 1
        elif status.startswith("SKIP already"):
            counts["SKIP already cypher"] += 1
        else:
            counts["other"] += 1
            failures.append((str(p), status))
        print(f"{status}\t{p}")

    print("\nSummary:", counts, "files:", len(uniq))
    print("Missing TC inputs:", len(missing))
    for m in missing:
        print("  MISSING", m)
    if failures:
        print("Failures:")
        for f in failures:
            print(" ", f)


if __name__ == "__main__":
    main()
